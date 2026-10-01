"""PDD Stage 8: zip the per-item files, write the reconciliation + execution
summary as an .xlsx, and (later) hand it to the email step.
"""
from __future__ import annotations

import time
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


def make_zip(files: list[Path], zip_path: Path) -> Path:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in files:
            zf.write(f, arcname=f.name)
    return zip_path


def write_summary_xlsx(
    dest: Path,
    *,
    flow: str,
    run_id: str,
    started: float,
    finished: float,
    expected: int,
    succeeded: int,
    failed_items: list[str],
    reconciliation: dict,
    recon_columns: tuple[str, ...] = ("Bucket", "Month", "Sum of States", "All India", "Delta", "Match"),
    recon_row_keys: tuple[str, ...] = ("bucket", "month", "sum_of_states", "all_india", "delta", "match"),
) -> Path:
    wb = Workbook()

    ex = wb.active
    ex.title = "Execution Summary"
    for row in [
        ["Metric", "Value"],
        ["Flow", flow],
        ["Run ID", run_id],
        ["States Processed", succeeded],
        ["Expected Files", expected],
        ["Successful Downloads", succeeded],
        ["Failed Downloads", len(failed_items)],
        ["Start Time", datetime.fromtimestamp(started).isoformat(timespec="seconds")],
        ["End Time", datetime.fromtimestamp(finished).isoformat(timespec="seconds")],
        ["Duration (min)", round((finished - started) / 60, 1)],
        ["Data validation result", _recon_verdict(reconciliation, succeeded)],
    ]:
        ex.append(row)
    if succeeded == 0:
        ex.append([])
        ex.append(["NOTE", "0 report files were produced - check the state filter "
                          "and the run log."])
    if failed_items:
        ex.append([])
        ex.append(["Failed items"])
        for it in failed_items:
            ex.append([it])

    rc = wb.create_sheet("Reconciliation")
    rc.append(list(recon_columns))
    for r in reconciliation.get("rows", []):
        rc.append([
            ("YES" if r[k] else "NO") if k == "match" else r.get(k)
            for k in recon_row_keys
        ])
    if reconciliation.get("file_errors"):
        er = wb.create_sheet("File Errors")
        er.append(["File", "Error"])
        for line in reconciliation["file_errors"]:
            fname, _, msg = line.partition(": ")
            er.append([fname, msg])

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def timestamp_slug() -> str:
    return time.strftime("%Y%m%d-%H%M%S")


def _recon_verdict(reconciliation: dict, succeeded: int) -> str:
    if reconciliation.get("label"):  # explicit label from the flow, if given
        return reconciliation["label"]
    if succeeded == 0:
        return "N/A (no report files)"
    if reconciliation.get("skipped"):
        return "Skipped"
    if reconciliation.get("ok") is True:
        return "Success"
    if reconciliation.get("ok") is False:
        return "Failed (mismatch - see Reconciliation sheet)"
    return "N/A"
